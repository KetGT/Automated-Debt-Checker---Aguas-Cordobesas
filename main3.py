import os
import threading  # Clave para que la app no se congele mientras corre Playwright
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright, TimeoutError as PlaywrightTimeoutError

def procesar_excel(archivo_ruta, headless_mode):
    """Maneja el bucle de las UFs y escribe en el archivo Excel"""
    columna_uf = "D"
    columna_deuda = "E"
    fila_inicio = 2

    try:
        wb = load_workbook(archivo_ruta)
        ws = wb.active
    except Exception as e:
        print(f"Error al abrir el Excel: {e}")
        return

    print("Iniciando navegador...")
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=headless_mode)
        page = browser.new_page()

        max_rows = ws.max_row
        for fila in range(fila_inicio, max_rows + 1):
            uf = ws[f"{columna_uf}{fila}"].value

            if not uf:
                continue

            if not str(uf).isdigit():
                print(f"UF {uf} no es numérica, saltando...")
                continue

            uf = str(uf)
            print(f"Consultando UF {uf} ({fila}/{max_rows})...")

            try:
                # DEVUELVE "SI" o "NO"
                resultado = consultar_uf(uf, page)
            except Exception as e:
                print(f"Error consultando UF {uf}: {e}")
                continue

            # Escribir en las celdas
            ws[f"{columna_deuda}{fila}"] = resultado
            
        browser.close()

    try:
        wb.save(archivo_ruta)
        print("¡Actualización completa y archivo guardado!")
    except Exception as e:
        print(f"Error al guardar el archivo: {e}")

def consultar_uf(uf, page):
    deuda = "NO"

    page.goto("https://www.aguascordobesas.com.ar/espacioClientes/seccion/gestionDeuda")
    page.fill("#searchUf", uf)
    page.dispatch_event("#btn-searchUf", "click")

    try:
        page.get_by_role(
            "heading",
            name="Comprobantes vencidos",
            exact=True
        ).wait_for(state="visible", timeout=10_000)

        # Encontró el elemento de comprobantes vencidos, lo que indica que hay deuda
        print("Hay comprobantes vencidos")
        deuda = "SI"

    except PlaywrightTimeoutError:
        # No apareció el elemento de comprobante vencidos dentro del tiempo de espera
        print("No aparecieron comprobantes vencidos")
        deuda = "NO"

    # CORRECCIÓN: El return debe estar fuera del except para que devuelva el valor en ambos casos
    return deuda

# --- INTERFAZ GRÁFICA Y FUNCIÓN MAIN ---

def seleccionar_archivo():
    """Abre el explorador de archivos para elegir el Excel"""
    ruta = filedialog.askopenfilename(
        title="Seleccionar archivo Excel",
        filetypes=[("Archivos Excel", "*.xlsx *.xls")]
    )
    if ruta:
        entrada_ruta.delete(0, tk.END)
        entrada_ruta.insert(0, ruta)

def iniciar_proceso():
    """Valida los datos y lanza el hilo secundario"""
    ruta = entrada_ruta.get()
    if not ruta or not os.path.exists(ruta):
        messagebox.showerror("Error", "Por favor, selecciona un archivo Excel válido.")
        return

    modo_oculto = var_headless.get()
    
    # Deshabilitar el botón para no ejecutar dos veces al mismo tiempo
    btn_iniciar.config(state=tk.DISABLED)
    btn_iniciar.config(text="Procesando...")

    # Creamos el hilo para que la interfaz no se trabe
    hilo = threading.Thread(target=ejecutar_y_notificar, args=(ruta, modo_oculto))
    hilo.start()

def ejecutar_y_notificar(ruta, modo_oculto):
    """Ejecuta el scraper y avisa cuando termina"""
    try:
        procesar_excel(ruta, modo_oculto)
        messagebox.showinfo("Éxito", "El proceso ha terminado y el archivo se ha guardado.")
    except Exception as e:
        messagebox.showerror("Error", f"Ocurrió un error inesperado:\n{e}")
    finally:
        # Reactivar el botón al terminar
        btn_iniciar.config(state=tk.NORMAL)
        btn_iniciar.config(text="Iniciar Consulta")

# Punto de entrada del programa
if __name__ == "__main__":
    # Configuración de la ventana principal
    root = tk.Tk()
    root.title("Consultor de Deuda - Aguas Cordobesas")
    root.geometry("450x200")
    root.resizable(False, False)

    # Variables de Tkinter
    var_headless = tk.BooleanVar(value=True)

    # Elementos de la interfaz (Widgets)
    tk.Label(root, text="Archivo Excel:").pack(pady=(15, 5))
    
    frame_archivo = tk.Frame(root)
    frame_archivo.pack(padx=20, fill="x")
    
    entrada_ruta = tk.Entry(frame_archivo)
    entrada_ruta.pack(side="left", fill="x", expand=True, padx=(0, 10))
    
    tk.Button(frame_archivo, text="Buscar", command=seleccionar_archivo).pack(side="right")

    tk.Checkbutton(
        root, 
        text="Modo silencioso (No mostrar el navegador)", 
        variable=var_headless
    ).pack(pady=15)

    btn_iniciar = tk.Button(
        root, 
        text="Iniciar Consulta", 
        command=iniciar_proceso,
        bg="#4CAF50", 
        fg="white", 
        font=("Arial", 10, "bold")
    )
    btn_iniciar.pack(pady=10)

    # Bucle principal de la interfaz
    root.mainloop()